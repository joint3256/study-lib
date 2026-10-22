#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把提取结果 JSON 组装成规范的"干净原题" Word（A4，中文宋体、标题黑体、
英文与数字 Times New Roman，图片居中、选项两列、无任何手写内容）。

JSON 结构：
{
  "title": "试卷标题",
  "blocks": [
    {"type":"section","text":"一、选择题（本大题共12小题…）"},
    {"type":"q","no":"1.","text":"题干（　）","opts":["A. …","B. …","C. …","D. …"]},
    {"type":"q","no":"2.","text":"题干（　）"},
    {"type":"image","path":"/abs/figs/q02.png","width_cm":7.0},
    {"type":"opts","items":["A. …","B. …","C. …","D. …"]},
    {"type":"para","text":"（1）小问…"},
    {"type":"blank"}
  ]
}
block 类型：section 大题标题
           q 题目(no,text；opts 可省略，适合无图或选项紧跟题干)
           image 干净配图(path,width_cm)
           opts 独立两列选项块(items)，用于"题干→图→选项"的题
           para 普通段落/小问
           blank 空行。

用法：python3 make_clean_doc.py extract.json [out.docx]
"""
import sys
from pathlib import Path
from extract_schema import load_extract
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def style_run(run, size=11.5, bold=False, cn="宋体", en="Times New Roman"):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), cn)


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "none")
        borders.append(e)
    # OOXML order: w:tblW -> w:tblBorders -> w:tblLook (insert before tblLook)
    look = tbl_pr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(borders)
    else:
        tbl_pr.append(borders)


def render_opts(doc, items):
    rows = (len(items) + 1) // 2
    tb = doc.add_table(rows=rows, cols=2)
    no_borders(tb)
    for i, o in enumerate(items):
        cp = tb.cell(i // 2, i % 2).paragraphs[0]
        cp.paragraph_format.space_after = Pt(1)
        style_run(cp.add_run(o))


def main(json_path, out_path=None):
    data = load_extract(json_path)
    out = Path(out_path) if out_path else Path(json_path).with_suffix('.docx')
    inputs = [Path(json_path)] + [Path(block['path']) for block in data['blocks']
                                 if block['type'] == 'image']
    # Compare resolved paths and file identity: a .docx alias can still be an input.
    # Validate before constructing/saving the document; preserve normal regeneration.
    for source in inputs:
        if out.resolve() == source.resolve() or (out.exists() and out.samefile(source)):
            raise ValueError('输出文件不能覆盖输入 JSON 或配图：%s' % source)
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2.0), Cm(2.0)
    s.left_margin, s.right_margin = Cm(2.2), Cm(2.2)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(10)
    style_run(t.add_run(data.get("title", "")), size=16, bold=True, cn="黑体")

    for b in data.get("blocks", []):
        typ = b.get("type")
        if typ == "section":
            p = doc.add_paragraph()
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(4)
            style_run(p.add_run(b["text"]), size=13, bold=True, cn="黑体")
        elif typ == "q":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            if b.get("no"):
                style_run(p.add_run(b["no"] + " "), bold=True)
            style_run(p.add_run(b["text"]))
            if b.get("opts"):
                render_opts(doc, b["opts"])
        elif typ == "opts":
            render_opts(doc, b.get("items", []))
        elif typ == "image":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(2)
            p.add_run().add_picture(b["path"], width=Cm(b.get("width_cm", 12.0)))
        elif typ == "para":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.25
            if b.get("indent"):
                p.paragraph_format.left_indent = Cm(b["indent"])
            style_run(p.add_run(b["text"]))
        elif typ == "blank":
            doc.add_paragraph()

    doc.save(out)
    print("saved:", out)


if __name__ == "__main__":
    from console_output import configure_stdout
    configure_stdout()
    a = sys.argv[1:]
    if not a:
        sys.exit("usage: make_clean_doc.py extract.json [out.docx]")
    try:
        main(a[0], a[1] if len(a) > 1 else None)
    except (OSError, ValueError) as exc:
        sys.exit("生成失败：%s" % exc)
