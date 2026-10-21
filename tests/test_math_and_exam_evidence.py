import copy
from decimal import Decimal, localcontext
from fractions import Fraction
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('evidence', ROOT / 'scripts/check_exam_evidence.py')
evidence = importlib.util.module_from_spec(spec)
spec.loader.exec_module(evidence)


class MathExamples(unittest.TestCase):
    def test_ellipse_dot_product_and_discriminant(self):
        # 用精确有理数检查化简恒等式，避免只测一次代入。
        for k2 in (Fraction(0), Fraction(1, 7), Fraction(4, 11), Fraction(3), Fraction(10)):
            denom = 1 + 4*k2
            product_x = (12*k2-4)/denom
            # sqrt(3)*(x1+x2) = 24*k² / (1+4k²)
            product_y = k2*(product_x - 24*k2/denom + 3)
            self.assertEqual(product_x+product_y, (11*k2-4)/denom)
            self.assertEqual(192*k2*k2-4*denom*(12*k2-4), 16*(1+k2))
        k2=Fraction(4,11)
        self.assertEqual((11*k2-4)/(1+4*k2), 0)
        self.assertNotEqual(Fraction(3)-Fraction(1,4), 0)  # 竖直线排除

    def test_fold_answer_and_perpendicular_bisector(self):
        with localcontext() as ctx:
            ctx.prec = 50
            root7=Decimal(7).sqrt()
            x=(32-8*root7)/3
            self.assertGreater(x, 0)
            self.assertLess(x, 6)
            self.assertLess(abs(x*x-(6-x)**2-(8-2*root7)**2), Decimal('1e-46'))
            # A=(0,0), B=(8,0), B'=(2√7,6), E=(8,x)
            self.assertLess(abs((2*root7)**2+36-64), Decimal('1e-46'))

    def test_tangent_chord_contacts(self):
        with localcontext() as ctx:
            ctx.prec=50
            for sign in (1,-1):
                t=(1+sign*Decimal(5).sqrt())/2
                x,y=t*t,2*t
                self.assertLess(abs(y*y-4*x),Decimal('1e-46'))
                self.assertLess(abs(y-(2*x-2)),Decimal('1e-46'))
                self.assertLess(abs(t-(-1+t*t)),Decimal('1e-46'))

    def test_old_circle_claim_counterexample(self):
        # A=(0,0),B=(4,-4),F=(1,0)：F 不在直径圆上。
        self.assertEqual((-1)*3+0*(-4),-3)

    def test_cube_line_intersects_side_plane(self):
        # A=(0,0,0),B=(2,0,0),C=(2,2,0),D=(0,2,0),E=(1,0,0)
        # DE 的参数式 (t,2-2t,0)，t=2 时穿过平面 x=2。
        t=2
        self.assertEqual((t,2-2*t,0),(2,-2,0))
        self.assertNotEqual((1,-2,0)[0],0)  # 与平面法向量 (1,0,0) 点积非零

    def test_recovered_rhombus_answer_by_coordinates(self):
        # 题面独立坐标复算；y 坐标以 sqrt(3) 为单位记录。
        a=(Fraction(0),Fraction(0)); f=(Fraction(2),Fraction(1,2))
        g=(Fraction(1),Fraction(1,4))
        self.assertEqual(g,tuple((x+y)/2 for x,y in zip(a,f)))
        self.assertEqual((f[0]-g[0])**2+3*(f[1]-g[1])**2,Fraction(19,16))

    def test_cli_error_on_non_utf8_console(self):
        with tempfile.TemporaryDirectory() as missing:
            result=subprocess.run([sys.executable, str(ROOT/'scripts/check_exam_evidence.py'),missing],
                                  env=dict(os.environ,PYTHONIOENCODING='ascii:strict'),capture_output=True)
        self.assertNotEqual(result.returncode,0)
        self.assertIn('证据检查失败',json.loads(result.stderr)['错误'])
        self.assertNotIn(b'UnicodeEncodeError',result.stderr)

    def test_corrected_budget_arithmetic(self):
        self.assertEqual(sum((24,12,16,8)),60)
        self.assertEqual(sum((21,3,12,7,3,8)),54)
        self.assertEqual(sum((14,13,6,3,2,4)),42)
        self.assertEqual(Fraction(sum((8,8,10,8,8)),5),Fraction(42,5))


class EvidenceTests(unittest.TestCase):
    def test_cli_json_on_non_utf8_console(self):
        # 重现 Windows 默认输出代码页无法表示中文时的 CLI 失败。
        env = dict(os.environ, PYTHONIOENCODING='ascii:strict', PYTHONUTF8='0')
        result = subprocess.run([sys.executable, '-B', str(ROOT / 'scripts/check_exam_evidence.py')],
                                env=env, capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report['候选记录总数'], 380)
        self.assertEqual(report['已核实记录'], 13)

    def record(self, **changes):
        r=dict(subject='数学',year=2022,question=1,topic='有理数',status='已核实',
               evidence='题面 https://example.org/paper.pdf 第1页 第1题；答案 https://example.org/key.pdf 第2页 第1题；来源 第三方原卷转载；核验 2026-10-21')
        r.update(changes)
        if 'evidence' not in changes:
            r['evidence'] = r['evidence'].replace('第1题', '第%d题' % r['question'])
        return r

    def test_pending_never_counts_as_occurrence(self):
        report=evidence.summarize([self.record(),self.record(question=2,status='待核实',evidence='—'),
                                  self.record(question=3,status='推断待核实',evidence='—')])
        self.assertEqual(report['已核实记录'],1)
        self.assertEqual(report['待核实记录'],2)
        self.assertEqual(report['已核实记录内计数'][0]['题目次数'],1)

    def test_year_count_differs_from_question_count(self):
        report=evidence.summarize([self.record(),self.record(question=2),self.record(year=2023)])
        entry=report['已核实记录内计数'][0]
        self.assertEqual(entry['出现年份数'],2)
        self.assertEqual(entry['题目次数'],3)

    def test_missing_evidence_and_inferred_cannot_pass(self):
        for changes in ({'evidence':'—'}, {'evidence':'https://example.org'},
                        {'evidence':'题面 https://example.org 第1页 2026-10-21'},
                        {'topic':'待核实（答案反推）'}, {'status':'猜测完成'},
                        {'evidence':'题面与答案 https://example.org 第1页 2026-99-99'}):
            with self.subTest(changes=changes),self.assertRaises(ValueError):
                evidence.summarize([self.record(**changes)])

    def test_answer_absence_cannot_be_verified(self):
        for invalid in (
            '题面 https://example.org/paper.pdf 第1页；答案未找到；2026-10-21',
            '题面与答案 https://example.org/paper.pdf 第1页 第1题；2026-10-21',
            '题面 https://example.org/paper.pdf 第1页 第1题；答案 https://example.org/key.pdf；2026-10-21',
            self.record()['evidence'].replace('第2页 第1题', '第2页 第2题'),
            self.record()['evidence'].replace('第2页', '第0页'),
            self.record()['evidence'].replace('https://example.org/key.pdf', 'https:///key.pdf'),
            self.record()['evidence'].replace('2026-10-21', '2099-10-21'),
            self.record()['evidence'].replace('2026-10-21', '2021-10-21'),
        ):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                evidence.summarize([self.record(evidence=invalid)])

    def test_separate_locations_can_share_source(self):
        r = self.record()
        r['evidence'] = r['evidence'].replace('key.pdf', 'paper.pdf')
        self.assertEqual(evidence.summarize([r])['已核实记录'], 1)

    def test_invalid_record_identity_rejected(self):
        for changes in ({'year':2027}, {'year':True}, {'subject':'未知'}, {'question':0}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                evidence.summarize([self.record(**changes)])

    def parse_mutated(self, transform):
        with tempfile.TemporaryDirectory() as temp:
            for subject in evidence.SUBJECTS:
                text = (ROOT/'真题考点/天津'/(subject+'.md')).read_text(encoding='utf-8')
                before, section = text.split('## 四、逐题对照明细', 1)
                (Path(temp)/(subject+'.md')).write_text(
                    before+'## 四、逐题对照明细'+transform(section), encoding='utf-8')
            return evidence.read_records(temp)

    def test_markdown_spacing_does_not_drop_records(self):
        rows = self.parse_mutated(lambda s: s.replace('| 1 |', '|1|', 1))
        self.assertEqual(len(rows), 380)

    def test_malformed_rows_and_years_fail_closed(self):
        for transform in (
            lambda s: s.replace('| 1 |', '| X |', 1),
            lambda s: s.replace('| 1 |', '1 |', 1),
            lambda s: s.replace('### 2023 年', '### 2027 年'),
            lambda s: s.replace('### 2023 年', '### 2022 年'),
            lambda s: s.replace('### 2023 年', '### 2023年（未核实）'),
            lambda s: s[:s.index('### 2026 年')] + s[s.index('## 五、'):],
        ):
            with self.subTest(transform=transform), self.assertRaises(ValueError):
                self.parse_mutated(transform)

    def test_duplicate_rows_rejected(self):
        with self.assertRaises(ValueError):
            evidence.summarize([self.record(),self.record()])

    def test_repository_known_gaps_remain_excluded(self):
        records=evidence.read_records(ROOT/'真题考点/天津')
        result=evidence.summarize(records)
        self.assertEqual(result['候选记录总数'],380)
        self.assertEqual(result['已核实记录'],13)
        self.assertEqual(result['待核实记录'],367)
        self.assertEqual(result['其中原推断记录'],0)
        self.assertEqual(result['已核实来源类型'],{'第三方解析版':3, '第三方原卷转载':10})
        verified={(r['subject'],r['year'],r['question']) for r in records if r['status']=='已核实'}
        self.assertEqual(verified,{('数学',2022,n) for n in (12,14,17)}|{('化学',2024,n) for n in range(1,11)})
        self.assertEqual(sum(x['候选记录'] for x in result['候选记录核验进度']),380)
        for subject in evidence.SUBJECTS:
            text=(ROOT/'真题考点/天津'/(subject+'.md')).read_text(encoding='utf-8')
            self.assertIn('暂停发布确定频次',text)
            self.assertNotIn('严格依据天津招考资讯网官方',text)
            self.assertNotIn('✓(',text)


class WordSyncTests(unittest.TestCase):
    def test_word_retains_all_pending_records_without_old_frequencies(self):
        from docx import Document
        doc=Document(ROOT/'真题考点/天津/天津中考数理化真题考点频次图谱_A层_2022-2026.docx')
        records=[r for table in doc.tables if len(table.columns)==7 for r in table.rows[1:]]
        self.assertEqual(len(records),380)
        md_records=evidence.read_records(ROOT/'真题考点/天津')
        for row, expected in zip(records,md_records):
            self.assertEqual([c.text for c in row.cells[:4]],
                             [str(expected['question']),expected['topic'],expected['question_type'],expected['score']])
            self.assertEqual(row.cells[5].text,expected['status'])
            self.assertEqual(row.cells[6].text,expected['evidence'])
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertNotIn('严格依据天津招考资讯网官方',text)
        self.assertNotIn('5/5',text)

    def test_parent_guides_withdraw_score_promises(self):
        from docx import Document
        for subject in evidence.SUBJECTS:
            path=ROOT/'真题考点/天津/家长版'/('天津中考'+subject+'_给家长的备考说明_2022-2026.docx')
            text='\n'.join(p.text for p in Document(path).paragraphs)
            self.assertIn('已撤回',text)
            self.assertNotIn('稳拿',text)
            self.assertNotIn('年年',text)


if __name__=='__main__':
    unittest.main()
