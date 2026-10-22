"""可见公式与隐藏属性回归；只使用临时合成文档。"""
import contextlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "技能/tiqu-extractor/scripts"
sys.path.insert(0, str(SCRIPTS))
import make_clean_doc
import verify_docx


def add_equation(paragraph):
    equation = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "answer=2"
    run.append(text)
    equation.append(run)
    paragraph._p.append(equation)


def set_switch(properties, name, value):
    node = OxmlElement("w:" + name)
    if value is not None:
        node.set(qn("w:val"), value)
    properties.append(node)


class DocxContentGuards(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.source = self.base / "input.json"
        self.target = self.base / "output.docx"
        self.source.write_text(json.dumps({"title": "Audit fixture", "blocks": [
            {"type": "q", "no": 1, "text": "Calculate 1+1.", "opts": ["A. 1", "B. 2"]}
        ]}), encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()):
            make_clean_doc.main(str(self.source), str(self.target))
        self.original = self.target.read_bytes()

    def document(self):
        return Document(io.BytesIO(self.original))

    def verify(self, doc):
        doc.save(self.target)
        return verify_docx.verify(str(self.source), str(self.target))

    def test_native_equations_are_not_silently_ignored(self):
        for location in ("paragraph", "option", "header", "footer"):
            with self.subTest(location=location):
                doc = self.document()
                paragraph = (doc.paragraphs[1] if location == "paragraph" else
                             doc.tables[0].cell(0, 0).paragraphs[0] if location == "option" else
                             getattr(doc.sections[0], location).paragraphs[0])
                add_equation(paragraph)
                report = self.verify(doc)
                self.assertFalse(report["ok"], report)
                self.assertIn("公式", "\n".join(report["errors"]))

    def test_explicit_hidden_off_is_accepted(self):
        for name in ("vanish", "webHidden"):
            for value in ("0", "false", "off"):
                with self.subTest(name=name, value=value):
                    doc = self.document()
                    set_switch(doc.paragraphs[1].runs[0]._r.get_or_add_rPr(), name, value)
                    self.assertTrue(self.verify(doc)["ok"])

    def test_enabled_and_invalid_hidden_values_are_rejected(self):
        for name in ("vanish", "webHidden"):
            for value in (None, "1", "true", "on", "not-a-boolean"):
                with self.subTest(name=name, value=value):
                    doc = self.document()
                    set_switch(doc.tables[0].cell(0, 0).paragraphs[0].runs[0]._r.get_or_add_rPr(),
                               name, value)
                    self.assertFalse(self.verify(doc)["ok"])

    def test_applied_style_hidden_text_is_rejected(self):
        for kind in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            with self.subTest(kind=kind):
                doc = self.document()
                style = doc.styles.add_style("AuditHidden", kind)
                style.font.hidden = True
                paragraph = doc.tables[0].cell(0, 0).paragraphs[0]
                if kind == WD_STYLE_TYPE.PARAGRAPH:
                    paragraph.style = style
                else:
                    paragraph.runs[0].style = style
                self.assertFalse(self.verify(doc)["ok"])

    def test_direct_off_overrides_hidden_style(self):
        doc = self.document()
        style = doc.styles.add_style("AuditHidden", WD_STYLE_TYPE.CHARACTER)
        style.font.hidden = True
        run = doc.paragraphs[1].runs[0]
        run.style = style
        run.font.hidden = False
        self.assertTrue(self.verify(doc)["ok"])

    def test_inherited_style_toggle_can_restore_visible_text(self):
        doc = self.document()
        hidden = doc.styles.add_style("AuditHidden", WD_STYLE_TYPE.CHARACTER)
        hidden.font.hidden = True
        visible = doc.styles.add_style("AuditVisible", WD_STYLE_TYPE.CHARACTER)
        visible.base_style = hidden
        visible.font.hidden = True
        doc.paragraphs[1].runs[0].style = visible
        self.assertTrue(self.verify(doc)["ok"])

    def test_default_hidden_text_is_rejected(self):
        doc = self.document()
        rpr = doc.styles.element.xpath('./w:docDefaults/w:rPrDefault/w:rPr')[0]
        set_switch(rpr, "vanish", "true")
        self.assertFalse(self.verify(doc)["ok"])

    def test_cli_rejects_equation_and_accepts_explicit_off(self):
        for equation, expected in ((True, 1), (False, 0)):
            with self.subTest(equation=equation):
                doc = self.document()
                if equation:
                    add_equation(doc.paragraphs[1])
                else:
                    doc.paragraphs[1].runs[0].font.hidden = False
                doc.save(self.target)
                result = subprocess.run([sys.executable, "-B", str(SCRIPTS / "verify_docx.py"),
                                         str(self.source), str(self.target)], capture_output=True)
                self.assertEqual(result.returncode, expected, result.stderr)
                self.assertEqual(json.loads(result.stdout)["ok"], expected == 0)


if __name__ == "__main__":
    unittest.main()
