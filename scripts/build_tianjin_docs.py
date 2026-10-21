#!/usr/bin/env python3
"""从已修订图谱重建仓库中的四份 Word，防止旧统计继续随成品流传。

依赖 python-docx；从仓库根目录运行 python scripts/build_tianjin_docs.py。
生成后仍须转 PDF 逐页检查；不会下载或嵌入原卷、教材。
"""
from pathlib import Path
from check_exam_evidence import read_records, summarize
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parents[1] / '真题考点/天津'
FONT = 'Noto Sans CJK SC'


def plain(value):
    value = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1（\2）', value)
    return value.replace('**', '').replace('`', '').replace('\\|', '|')


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.left_margin = section.right_margin = Cm(1.5)
    for name, size in [('Normal', 10), ('Title', 20), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.12
    # 删除模板继承的标题横线，避免生成器默认装饰覆盖标题样式。
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    return doc


def table(doc, rows):
    count = len(rows[0])
    widths = {3:[1.3,8.6,8.1], 4:[1.1,7.1,6.1,3.7], 7:[0.8,7.2,1.3,1.2,1.4,2.5,3.6]}.get(count, [18/count]*count)
    if count == 3 and rows[0][0] == '学习内容':
        widths = [4,7,7]
    tb = doc.add_table(rows=0, cols=count)
    tb.autofit = False
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(tb.columns, widths):
        col.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        border=OxmlElement('w:'+side)
        border.set(qn('w:val'),'single'); border.set(qn('w:sz'),'4'); border.set(qn('w:color'),'D9D9D9')
        borders.append(border)
    tb._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        assert len(values) == count, values
        row=tb.add_row()
        trpr=row._tr.get_or_add_trPr()
        if index==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        # 单行跨页会让状态与题目分离，保持每条记录完整。
        trpr.append(OxmlElement('w:cantSplit'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width=Cm(width)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margin=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'70'); e.set(qn('w:type'),'dxa'); margin.append(e)
            tcpr.append(margin)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF5' if index==0 else ('F7F8FA' if index%2==0 else 'FFFFFF')); tcpr.append(shade)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.05
            r=p.add_run(plain(value))
            r.bold=index==0
            r.font.size=Pt(9 if count>=4 else 9.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return tb


def append_markdown(doc, text):
    lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line or line=='---': i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=lines[i].strip()
                if not re.fullmatch(r'[| :\-]+',row):
                    rows.append([c.strip() for c in re.split(r'(?<!\\)\|', row.strip('|'))])
                i+=1
            table(doc, rows)
            continue
        heading=re.match(r'^(#{1,3})\s+(.+)',line)
        if heading:
            title=re.sub(r'[、，。：“”（）()—·]+',' ',heading[2]).strip()
            doc.add_heading(title,level=min(len(heading[1]),3))
        else:
            if line.startswith('> '): line=line[2:]
            if line.startswith('- '): line='• '+line[2:]
            doc.add_paragraph(plain(line))
        i+=1


BASICS={
'数学': [('有理数 整式 分式','计算与化简','检查符号、运算顺序和分母限制'),('方程 不等式 函数','列关系式与读图','说清每一步变形依据，检验结果'),('几何与三角函数','识别条件、画辅助线','说明全等、相似或勾股定理的适用条件'),('统计与概率','读表、算统计量或概率','核对总数、单位和分类是否完整')],
'物理': [('声 光 热与生活常识','概念辨析与现象解释','先说物理过程，再选对应规律'),('力学与电学基础','受力图、电路图和单位换算','检查对象、方向、连接方式及单位'),('实验与测量','读数、作图和控制变量','区分操作、观察结果与结论'),('热量与效率','列已知量、用条件选公式','核对单位、能量去向和结果范围')],
'化学': [('化学用语与方程式','元素符号、化合价和配平','检查原子数守恒、状态及反应条件'),('常见物质性质','性质、用途与分类','用反应或现象说明判断依据'),('实验操作与气体制取','装置选择和操作次序','说明选择依据，区分现象与结论'),('化学式与溶液计算','质量比、质量分数及守恒','列关系式、标单位、核对结果范围')]
}
COMPLEX={
'数学':'圆的证明、图形折叠与平移、函数综合、动点与分类讨论，可以先练单一模型，再练多个条件的结合。不要根据旧图谱认定某题号每年固定考某个知识点。',
'物理':'浮力、密度测量和电路设计可以先拆成受力、浮沉条件、测量关系或电路连接等小问题。多选题先阅读当次试卷的评分规则，再判断每个选项。',
'化学':'酸碱盐推断、物质转化、金属活动性和综合计算，可以先列出已知现象与反应，再逐步建立数量关系。实验题的分值与题号必须对照具体试卷。'
}
GAPS={
'数学':'2022年25题已逐题核对第三方材料。第18题是等腰直角三角形的网格构造，原“切线”标签有误；第20题没有用样本估计总体的小问，已修正。其余四年100条仍待核实，不能恢复五年频次。旧“基础盘约80分”列项实际合计60分，也不应作为孩子的保证得分。',
'物理':'旧资料未保存足以复查的逐题题面与答案定位，2026年部分综合题表达式也留有疑问。旧“基础盘约60分”列项合计54分，且可能重复计入小问，因此撤销这个目标。',
'化学':'2024年26题已逐题核对第三方转载的扫描卷与参考答案；第11至15题为有1或2个正确项的不定项选择。该年完成核验不代表其他四年完成，仍不能恢复五年频次。旧“基础盘约60分”列项实际合计42分，后续分组又重复计入部分题目，因此撤销60/30/10分分组。'
}


def build_parent(subject):
    doc=new_document()
    doc.styles['Normal'].paragraph_format.space_after=Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing=1.05
    doc.styles['Heading 1'].font.size=Pt(14)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
    doc.styles['Title'].font.size=Pt(18)
    doc.add_paragraph('天津中考'+subject+'给家长的备考说明',style='Title')
    doc.add_paragraph('2022至2026年资料核验修订  2026年10月21日')
    doc.add_paragraph('这份说明保留基础练习建议。原资料中的五年频次、固定题号和“必拿分”结论缺少完整逐题证据，已撤回。请根据孩子实际错题安排学习，不把旧统计当作考试预测或得分保证。')
    doc.add_heading('一 先确认考试信息',level=1)
    doc.add_paragraph('考试时长、题量、分值和多选题计分规则，请核对对应年份的正式试卷或考试说明。正在核验的旧记录不能替代这些信息。')
    doc.add_heading('二 先练基础知识',level=1)
    table(doc,[['学习内容','练习重点','家长可检查的表现']]+[list(r) for r in BASICS[subject]])
    doc.add_heading('三 综合题逐步练',level=1)
    doc.add_paragraph(COMPLEX[subject])
    doc.add_paragraph('先让孩子独立读题、说出条件，完成能解释依据的小问；遇到卡点，再回到对应基础知识。练习后检查中间步骤，不只看最后答案。')
    doc.add_heading('四 历年规律仍待核实',level=1)
    doc.add_paragraph(GAPS[subject])
    doc.add_paragraph('目前保留的候选记录，不等于已确认的真题事实。以后只有补齐题面、答案、来源定位和核验日期，才能重新统计频次；“尚未核实”也不代表这个考点没有考过。')
    doc.add_heading('五 按错题安排练习',level=1)
    doc.add_paragraph('先选少量近期错题，区分是概念不清、计算错误、读题遗漏还是步骤表达不完整。每次集中解决一个问题，再用同类题检查是否会独立完成。基础稳定后逐步增加综合度，练习目标由实际表现调整。')
    doc.save(BASE/'家长版'/('天津中考'+subject+'_给家长的备考说明_2022-2026.docx'))


def build_atlas():
    doc=new_document()
    doc.add_paragraph('天津中考数理化真题考点核验图谱',style='Title')
    doc.add_paragraph('2022至2026年  证据状态修订  2026年10月21日')
    report = summarize(read_records(BASE))
    doc.add_paragraph('本版保存%d条候选记录，已核实%d条，待核实%d条。已核实条目仍需区分官方原卷与第三方转载，详见各条证据。已核验范围按问题优先级选取，并非随机抽样；不得用这部分的结果推断其余记录的准确率。暂不发布五年频次、必考排序和年均分值。' % (report['候选记录总数'], report['已核实记录'], report['待核实记录']))
    for index,subject in enumerate(('数学','物理','化学')):
        start = len(doc.paragraphs)
        append_markdown(doc,(BASE/(subject+'.md')).read_text(encoding='utf-8'))
        if index:
            # 换学科时让标题自身分页，避免前一页刚好满时独立分页段落产生空白页。
            doc.paragraphs[start].paragraph_format.page_break_before = True
    doc.save(BASE/'天津中考数理化真题考点频次图谱_A层_2022-2026.docx')


def main():
    build_atlas()
    for subject in ('数学','物理','化学'): build_parent(subject)
    print('已更新图谱及三份家长版；请继续执行逐页渲染核对。')


if __name__=='__main__':
    main()
