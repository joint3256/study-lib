import hashlib
import importlib.util
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
EXTRACT = ROOT / "技能/tiqu-extractor/scripts"
sys.path.insert(0, str(EXTRACT))
import make_clean_doc
import verify_docx

spec = importlib.util.spec_from_file_location("repo_links", ROOT / "scripts/check_repository_links.py")
links = importlib.util.module_from_spec(spec)
spec.loader.exec_module(links)


class WordAuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.image = self.base / "图 中文.png"
        Image.new("RGB", (100, 80), "white").save(self.image)
        self.source = self.base / "extract.json"
        self.data = {"title": "合成测试卷", "blocks": [
            {"type": "section", "text": "一、选择题"},
            {"type": "q", "no": 1, "text": "请选择（　）", "opts": ["A. 甲", "B. 乙", "C. 丙"]},
            {"type": "q", "no": "2.", "text": "观察图形"},
            {"type": "image", "path": self.image.name, "width_cm": 5},
            {"type": "opts", "items": ["A. 正方形", "B. 三角形"]},
            {"type": "para", "text": "（1）填写理由：______"}, {"type": "blank"}]}
        self.source.write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8")
        self.target = self.base / "output.docx"
        make_clean_doc.main(str(self.source), str(self.target))

    def verify(self):
        return verify_docx.verify(str(self.source), str(self.target))

    def test_valid_document_and_read_only(self):
        before = {p.name: (hashlib.sha256(p.read_bytes()).hexdigest(), p.stat().st_mtime_ns)
                  for p in self.base.iterdir()}
        self.assertTrue(self.verify()["ok"])
        self.assertEqual(before, {p.name: (hashlib.sha256(p.read_bytes()).hexdigest(), p.stat().st_mtime_ns)
                                 for p in self.base.iterdir()})

    def test_modified_question_and_missing_paragraph_fail(self):
        doc = Document(self.target)
        doc.paragraphs[2].text = "1. 错误题干"
        doc.save(self.target)
        self.assertFalse(self.verify()["ok"])
        element = doc.paragraphs[2]._element
        element.getparent().remove(element)
        doc.save(self.target)
        self.assertFalse(self.verify()["ok"])

    def test_missing_or_reordered_options_fail(self):
        doc = Document(self.target)
        doc.tables[0].cell(0, 0).text = "B. 乙"
        doc.tables[0].cell(0, 1).text = "A. 甲"
        doc.save(self.target)
        self.assertFalse(self.verify()["ok"])
        doc.tables[0].cell(0, 0).text = ""
        doc.save(self.target)
        self.assertFalse(self.verify()["ok"])

    def test_image_content_changed_fails(self):
        Image.new("RGB", (100, 80), "red").save(self.image)
        self.assertFalse(self.verify()["ok"])

    def test_missing_image_and_oversized_image_fail(self):
        doc = Document(self.target)
        doc.inline_shapes[0].width = Cm(30)
        doc.save(self.target)
        self.assertFalse(self.verify()["ok"])
        element = doc.inline_shapes[0]._inline
        element.getparent().remove(element)
        doc.save(self.target)
        self.assertFalse(self.verify()["ok"])

    def test_broken_or_missing_docx_fails(self):
        self.target.write_bytes(b"not docx")
        self.assertFalse(self.verify()["ok"])
        self.target.unlink()
        self.assertFalse(self.verify()["ok"])

    def test_unexpected_footer_and_hidden_text_fail(self):
        doc = Document(self.target)
        doc.sections[0].footer.paragraphs[0].text = "原卷没有的答案"
        doc.save(self.target)
        self.assertFalse(self.verify()["ok"])
        doc.sections[0].footer.paragraphs[0].text = ""
        doc.paragraphs[2].runs[0]._r.get_or_add_rPr().append(OxmlElement('w:vanish'))
        doc.save(self.target)
        self.assertFalse(self.verify()["ok"])

    def test_cli_from_other_working_directory(self):
        for code in (0, 1):
            if code:
                self.target.write_bytes(b"broken")
            result = subprocess.run([sys.executable, "-B", str(EXTRACT / "verify_docx.py"),
                                     str(self.source), str(self.target)], cwd=self.base.parent,
                                    capture_output=True)
            self.assertEqual(result.returncode, code, result.stderr)
            self.assertEqual(json.loads(result.stdout)["ok"], code == 0)


class RepositoryLinksTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name).resolve()
        self.doc = self.base / "README.md"

    def test_relative_unicode_link_and_url_encoded_path(self):
        (self.base / "目录 中文").mkdir()
        (self.base / "目录 中文/目标.md").write_text("# 内容", encoding="utf-8")
        self.doc.write_text('[入口](目录%20中文/目标.md#未检查锚点)\n[官网](https://example.org/)', encoding="utf-8")
        self.assertEqual(links.local_link_errors(self.base), [])

    def test_missing_duplicated_directory_rejected(self):
        self.doc.write_text('[专题](初中/初中/01.md)', encoding="utf-8")
        self.assertIn('链接目标缺失', links.local_link_errors(self.base)[0])

    def test_stale_absolute_paths_rejected(self):
        self.doc.write_text('[文件](C:\\旧目录\\file.md)\n`/home/user/Doubao/chats/123/lib`', encoding="utf-8")
        self.assertEqual(len(links.local_link_errors(self.base)), 2)

    def test_code_fence_examples_not_real_links(self):
        self.doc.write_text('```md\n[示例](不存在.md)\n```\n~~~md\n[示例](不存在2.md)\n~~~', encoding="utf-8")
        self.assertEqual(links.local_link_errors(self.base), [])

    def test_precise_root_relative_file_checked(self):
        self.doc.write_text('文件：`方法套路/缺失.md`；`方法套路/<学科>/*.md`', encoding="utf-8")
        self.assertEqual(len(links.local_link_errors(self.base)), 1)

    def test_missing_root_is_failure(self):
        self.assertFalse(links.check(self.base / "missing")["ok"])

    def test_real_repository_links_indexes_and_counts(self):
        report = links.check(ROOT)
        self.assertTrue(report["ok"], report)
        self.assertIn(f'{report["method_total"]} 个方法/模型', (ROOT / 'README.md').read_text(encoding='utf-8'))
        self.assertIn(f'{report["method_total"]} 个方法/模型', (ROOT / '方法套路/README.md').read_text(encoding='utf-8'))
        self.assertEqual(report['method_counts']['初中数学'], 36)
        self.assertEqual(report['method_counts']['物理'], 70)

    def test_eight_sections_unchanged(self):
        files = [ROOT / '技能/cz-combine/SKILL.md'] + sorted((ROOT / '技能/cz-combine/references').glob('*专精.md'))
        for path in files:
            text = path.read_text(encoding='utf-8')
            rows = re.findall(r'^\| ([1-8]) \| ([^|]+)\| ([^|]+)\|', text, re.M)
            # Only numbered teaching rows carry the shared palette; other numbered examples may coexist.
            rows = [row for row in rows if any(c in row[2] for c in ('F4F7FB','DEEBF7','E2EFDA','222222','FCE4D6','FFF2CC','FDECEA','E4DFEC'))]
            self.assertEqual([row[0] for row in rows], list('12345678'), str(path))
            self.assertIn('互动对白', rows[-1][1])


if __name__ == '__main__':
    unittest.main()
