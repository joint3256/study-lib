"""复制自不同平台的多行题目应保持换行数量，并通过生成后核对。"""
import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '技能/tiqu-extractor/scripts'))
from extract_schema import validate_extract
from make_clean_doc import main as render
from verify_docx import verify


class NewlineTests(unittest.TestCase):
    def test_multiline_fields_round_trip_without_extra_line_breaks(self):
        for newline in ('\n', '\r\n', '\r'):
            with self.subTest(newline=repr(newline)), tempfile.TemporaryDirectory() as td:
                root = Path(td)
                value = '第一行' + newline + '第二行'
                data = {'title': value, 'blocks': [
                    {'type': 'section', 'text': value},
                    {'type': 'q', 'no': 1, 'text': value, 'opts': [value, 'B. 其他']},
                    {'type': 'para', 'text': value},
                    {'type': 'opts', 'items': [value, 'D. 其他']},
                ]}
                source, output = root / 'extract.json', root / 'out.docx'
                source.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
                original = source.read_bytes()
                with contextlib.redirect_stdout(io.StringIO()):
                    render(source, output)
                report = verify(source, output)
                self.assertTrue(report['ok'], report)
                doc = Document(output)
                self.assertEqual([p.text for p in doc.paragraphs], [
                    '第一行\n第二行', '第一行\n第二行',
                    '1. 第一行\n第二行', '第一行\n第二行',
                ])
                for table in doc.tables:
                    self.assertEqual(table.cell(0, 0).text, '第一行\n第二行')
                self.assertEqual(source.read_bytes(), original)

    def test_blank_lines_tabs_and_source_object_are_preserved(self):
        text = '题干\r\n\r\n\t小问\r末行\n'
        data = {'title': '试卷', 'blocks': [{'type': 'q', 'no': 1, 'text': text}]}
        normalized = validate_extract(data, Path('.'))
        self.assertEqual(normalized['blocks'][0]['text'], '题干\n\n\t小问\n末行\n')
        self.assertEqual(data['blocks'][0]['text'], text)


if __name__ == '__main__':
    unittest.main()
