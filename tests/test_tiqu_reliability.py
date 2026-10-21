"""临时文件回归；不读写用户的试卷或转换成果。"""
import contextlib
import copy
import importlib.util
import io
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / '技能/tiqu-extractor/scripts'
sys.path.insert(0, str(SCRIPTS))
import doc_to_outputs as convert
import make_clean_doc as render
import verify_extract as verify
from extract_schema import load_extract, validate_extract
from docx import Document
from PIL import Image

spec = importlib.util.spec_from_file_location('fig_helpers', ROOT / '技能/tiqu-extractor/assets/fig_helpers.py')
figures = importlib.util.module_from_spec(spec)
spec.loader.exec_module(figures)


class InputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root / '试卷.json'
        self.data = {'title': '试卷', 'blocks': [{'type': 'q', 'no': '1.', 'text': '计算1+1。'}]}

    def save(self, data):
        self.input.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')

    def assert_rejected(self, data):
        self.save(data)
        target = self.root / 'existing.docx'
        target.write_bytes(b'existing document')
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertNotEqual(verify.main(str(self.input)), 0)
        with self.assertRaises(ValueError):
            render.main(str(self.input), str(target))
        self.assertEqual(target.read_bytes(), b'existing document')

    def test_number_formats_work_in_both_tools(self):
        for value in (1, '1', '1.', '1．', '1、'):
            with self.subTest(value=value):
                self.data['blocks'][0]['no'] = value
                self.save(self.data)
                original = self.input.read_bytes()
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(verify.main(str(self.input)), 0)
                    render.main(str(self.input))
                text = '\n'.join(p.text for p in Document(self.input.with_suffix('.docx')).paragraphs)
                self.assertIn('1. 计算1+1。', text)
                self.assertEqual(self.input.read_bytes(), original)

    def test_invalid_root_and_blocks(self):
        for value in ([], None, {'title': 42, 'blocks': []},
                      {'title': 'a'}, {'title': 'a', 'blocks': {}},
                      {'title': 'a', 'blocks': [None]}, {'title': 'a', 'blocks': []}):
            with self.subTest(value=value):
                self.assert_rejected(value)

    def test_unknown_block_and_missing_text(self):
        for block in ({'type': 'qeustion', 'no': '1', 'text': '题目'},
                      {'type': 'q', 'no': 1}, {'type': 'q', 'no': 1, 'text': 123},
                      {'type': 'q', 'no': 1, 'text': '  '}, {'type': 'blank'}):
            with self.subTest(block=block):
                self.assert_rejected({'title': '试卷', 'blocks': [block]})

    def test_bad_number_types(self):
        for number in (True, False, 0, -1, 1.1, [], None, '第一题'):
            with self.subTest(number=number):
                self.data['blocks'][0]['no'] = number
                self.assert_rejected(self.data)

    def test_duplicate_number(self):
        self.data['blocks'].append({'type': 'q', 'no': 1, 'text': '重复题号'})
        self.assert_rejected(self.data)

    def test_options_types(self):
        for opts in (None, 'ABCD', [], [1, 2, 3, 4], ['A', '']):
            with self.subTest(opts=opts):
                self.data['blocks'][0]['opts'] = opts
                self.assert_rejected(self.data)

    def test_corrupt_missing_empty_and_directory_images(self):
        (self.root / 'corrupt.png').write_text('not an image')
        (self.root / 'empty.png').touch()
        for name in ('corrupt.png', 'missing.png', 'empty.png', '.'):
            with self.subTest(name=name):
                data = copy.deepcopy(self.data)
                data['blocks'].append({'type': 'image', 'path': name})
                self.assert_rejected(data)

    def test_unsupported_word_image(self):
        Image.new('RGB', (20, 20)).save(self.root / 'image.webp')
        self.data['blocks'].append({'type': 'image', 'path': 'image.webp'})
        self.assert_rejected(self.data)

    def test_image_relative_to_json_from_another_cwd(self):
        Image.new('RGB', (20, 30)).save(self.root / '图.png')
        self.data['blocks'].append({'type': 'image', 'path': '图.png', 'width_cm': 4})
        self.save(self.data)
        normalized = load_extract(self.input)
        normalized_image = Path(normalized['blocks'][1]['path'])
        self.assertTrue(normalized_image.is_absolute())
        # Windows 临时目录可使用 8.3 短名；核对文件身份，而非路径字符串的拼写。
        self.assertTrue(normalized_image.samefile(self.root / '图.png'))
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(verify.main(str(self.input)), 0)
            render.main(str(self.input))
        self.assertEqual(len(Document(self.input.with_suffix('.docx')).inline_shapes), 1)

    def test_invalid_image_width_and_indent(self):
        Image.new('RGB', (20, 20)).save(self.root / 'good.png')
        for value in (0, -1, True, float('inf'), float('nan'), '10'):
            with self.subTest(width=value):
                data = copy.deepcopy(self.data)
                data['blocks'].append({'type': 'image', 'path': 'good.png', 'width_cm': value})
                self.assert_rejected(data)
        data = copy.deepcopy(self.data)
        data['blocks'].append({'type': 'para', 'text': '小问', 'indent': -1})
        self.assert_rejected(data)

    def test_all_supported_blocks(self):
        self.data['blocks'] += [
            {'type': 'section', 'text': '第二部分'},
            {'type': 'opts', 'items': ['A.甲', 'B.乙', 'C.丙', 'D.丁']},
            {'type': 'para', 'text': '（1）说明理由', 'indent': 0}, {'type': 'blank'}]
        self.save(self.data)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(verify.main(str(self.input)), 0)
            render.main(str(self.input))
        self.assertEqual(len(Document(self.input.with_suffix('.docx')).tables), 1)

    def test_inline_options_get_same_space_check(self):
        opts = ['A. 6 米', 'B.7米', 'C.8米', 'D.9米']
        for inline in (True, False):
            data = copy.deepcopy(self.data)
            if inline:
                data['blocks'][0]['opts'] = opts
            else:
                data['blocks'].append({'type': 'opts', 'items': opts})
            self.save(data)
            report = io.StringIO()
            with contextlib.redirect_stdout(report):
                self.assertEqual(verify.main(str(self.input)), 0)
            self.assertIn('【CHECK 疑似（必须调出原照片逐处确认并记录）】共 1 处', report.getvalue())

    def test_missing_pillow_is_an_error(self):
        self.data['blocks'].append({'type': 'image', 'path': 'anything.png'})
        with patch.dict(sys.modules, {'PIL': None}):
            with self.assertRaisesRegex(ValueError, 'Pillow'):
                validate_extract(self.data, self.root)

    def test_invalid_json_and_cli(self):
        self.input.write_text('{invalid', encoding='utf-8')
        for script in ('verify_extract.py', 'make_clean_doc.py'):
            result = subprocess.run([sys.executable, '-B', str(SCRIPTS / script), str(self.input)],
                                    capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertNotIn('Traceback', result.stdout + result.stderr)


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / '试卷 空格.docx'
        Document().save(self.source)
        self.output = self.root / 'outputs'
        (self.output / 'pages').mkdir(parents=True)
        (self.output / 'pages/previous.png').write_bytes(b'previous image')
        (self.output / '_pdf').mkdir()
        (self.output / '_pdf/other.pdf').write_bytes(b'%PDF-old other')
        (self.output / '_pdf' / (self.source.stem + '.pdf')).write_bytes(b'%PDF-old same name')
        self.before = self.snapshot()

    def snapshot(self):
        return {str(p.relative_to(self.output)): p.read_bytes()
                for p in self.output.rglob('*') if p.is_file()}

    def fake_run(self, cmd):
        if cmd[0] == 'soffice':
            directory = Path(cmd[cmd.index('--outdir') + 1])
            (directory / (self.source.stem + '.pdf')).write_bytes(b'%PDF-new')
        else:
            for number in (1, 2, 10, 3, 4, 5, 6, 7, 8, 9):
                Image.new('RGB', (20, 20)).save(cmd[-1] + '-%d.png' % number)

    def test_success_preserves_old_files_and_sorts_numerically(self):
        with patch.object(convert, 'run', self.fake_run), contextlib.redirect_stdout(io.StringIO()):
            pdf, pages = convert.main(self.source, self.output)
            pdf2, _ = convert.main(self.source, self.output)
        self.assertNotEqual(pdf, pdf2)
        self.assertEqual([Path(p).name for p in pages], ['page-%d.png' % n for n in range(1, 11)])
        for name, contents in self.before.items():
            self.assertEqual((self.output / name).read_bytes(), contents)
        self.assertEqual(Path(pdf).read_bytes(), b'%PDF-new')

    def test_missing_pdf_does_not_reuse_stale_pdf(self):
        with patch.object(convert, 'run') as run:
            with self.assertRaisesRegex(RuntimeError, '预期文件'):
                convert.main(self.source, self.output)
        self.assertEqual(run.call_count, 1)
        self.assertEqual(self.snapshot(), self.before)

    def test_wrong_filename_is_rejected(self):
        def wrong(cmd):
            (Path(cmd[cmd.index('--outdir') + 1]) / 'other.pdf').write_bytes(b'%PDF-other')
        with patch.object(convert, 'run', wrong):
            with self.assertRaises(RuntimeError):
                convert.main(self.source, self.output)
        self.assertEqual(self.snapshot(), self.before)

    def test_converter_failure_preserves_previous_outputs(self):
        with patch.object(convert, 'run', side_effect=subprocess.CalledProcessError(1, ['soffice'])):
            with self.assertRaises(subprocess.CalledProcessError):
                convert.main(self.source, self.output)
        self.assertEqual(self.snapshot(), self.before)

    def test_missing_pages_are_a_failure(self):
        def no_pages(cmd):
            if cmd[0] == 'soffice':
                self.fake_run(cmd)
        with patch.object(convert, 'run', no_pages):
            with self.assertRaisesRegex(RuntimeError, 'PNG'):
                convert.main(self.source, self.output)
        self.assertEqual(self.snapshot(), self.before)

    def test_missing_input_does_not_create_outputs(self):
        with self.assertRaises(FileNotFoundError):
            convert.main(self.root / 'absent.docx', self.output)
        self.assertEqual(self.snapshot(), self.before)

    def test_text_disguised_as_docx_is_rejected(self):
        self.source.write_text('not a docx', encoding='utf-8')
        with patch.object(convert, 'run') as run:
            with self.assertRaisesRegex(ValueError, 'DOCX'):
                convert.main(self.source, self.output)
        run.assert_not_called()
        self.assertEqual(self.snapshot(), self.before)

    @unittest.skipUnless(shutil.which('soffice') and shutil.which('pdftoppm'), '需要 LibreOffice 和 pdftoppm')
    def test_real_conversion_twice_preserves_earlier_outputs(self):
        doc = Document()
        doc.add_paragraph('Regression document 1 + 1 = 2')
        doc.save(self.source)
        with contextlib.redirect_stdout(io.StringIO()):
            pdf, pages = convert.main(self.source, self.output, 72)
            earlier = self.snapshot()
            pdf2, pages2 = convert.main(self.source, self.output, 72)
        self.assertNotEqual(pdf, pdf2)
        self.assertTrue(pages and pages2)
        for name, contents in earlier.items():
            self.assertEqual((self.output / name).read_bytes(), contents)
        with Image.open(pages[0]) as image:
            image.verify()


class GeometryTests(unittest.TestCase):
    def test_minor_angles_and_reverse_order(self):
        for first, second, expected in ((170, -170, 20), (-170, 170, 20),
                                         (10, 80, 70), (80, 10, 70), (0, 180, 180)):
            with self.subTest(first=first, second=second):
                fig, ax = figures.newax((-2, 2), (-2, 2))
                q1 = (math.cos(math.radians(first)), math.sin(math.radians(first)))
                q2 = (math.cos(math.radians(second)), math.sin(math.radians(second)))
                figures.ang(ax, (0, 0), q1, q2)
                arc = ax.patches[-1]
                self.assertAlmostEqual(arc.theta2 - arc.theta1, expected)
                figures.plt.close(fig)

    def test_zero_length_ray_rejected(self):
        fig, ax = figures.newax((-1, 1), (-1, 1))
        with self.assertRaises(ValueError):
            figures.ang(ax, (0, 0), (0, 0), (1, 0))
        figures.plt.close(fig)


if __name__ == '__main__':
    unittest.main()
