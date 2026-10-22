"""有限字符集终端不能把成功的提取/转换误报为失败。"""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "技能/tiqu-extractor/scripts"


class ExtractorConsoleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.source = self.base / "二次函数x².json"
        self.data = {"title": "二次函数x²", "blocks": [
            {"type": "q", "no": 1, "text": "计算x²+1。"}]}
        self.source.write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8")

    def cli(self, script, *args, encoding="gbk"):
        return subprocess.run([sys.executable, "-B", str(SCRIPTS / script), *map(str, args)],
                              env=dict(os.environ, PYTHONIOENCODING=encoding + ":strict"),
                              capture_output=True)

    def test_generate_and_check_succeed_with_limited_console(self):
        original = self.source.read_bytes()
        for encoding in ("ascii", "gbk"):
            for script in ("make_clean_doc.py", "verify_extract.py"):
                with self.subTest(encoding=encoding, script=script):
                    result = self.cli(script, self.source, encoding=encoding)
                    self.assertEqual(result.returncode, 0, result.stderr)
                    self.assertNotIn(b"UnicodeEncodeError", result.stderr)
                    self.assertIn(b"\\xb2", result.stdout)
        self.assertEqual(self.source.read_bytes(), original)
        doc = Document(self.source.with_suffix(".docx"))
        self.assertEqual(doc.paragraphs[0].text, self.data["title"])
        self.assertEqual(doc.paragraphs[1].text, "1. " + self.data["blocks"][0]["text"])

    def test_utf8_output_keeps_literal_text(self):
        for script in ("make_clean_doc.py", "verify_extract.py"):
            with self.subTest(script=script):
                result = self.cli(script, self.source, encoding="utf-8")
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertIn("二次函数x²", result.stdout.decode("utf-8"))
                self.assertNotIn(b"\\xb2", result.stdout)

    def test_invalid_inputs_still_fail_cleanly(self):
        self.source.write_text("{invalid", encoding="utf-8")
        for encoding in ("ascii", "gbk"):
            for script, source in (("make_clean_doc.py", self.source),
                                   ("verify_extract.py", self.source),
                                   ("doc_to_outputs.py", self.base / "缺少x².docx")):
                with self.subTest(encoding=encoding, script=script):
                    result = self.cli(script, source, encoding=encoding)
                    self.assertNotEqual(result.returncode, 0)
                    output = result.stdout + result.stderr
                    self.assertNotIn(b"UnicodeEncodeError", output)
                    self.assertNotIn(b"Traceback", output)

    def test_imports_leave_stdout_policy_unchanged(self):
        code = ("import sys; sys.path.insert(0, %r); "
                "import make_clean_doc, verify_extract, doc_to_outputs; "
                "print(sys.stdout.encoding, sys.stdout.errors)") % str(SCRIPTS)
        result = subprocess.run([sys.executable, "-B", "-c", code],
                                env=dict(os.environ, PYTHONIOENCODING="gbk:strict"),
                                capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), b"gbk strict")

    @unittest.skipUnless(shutil.which("soffice") and shutil.which("pdftoppm"),
                         "需要 LibreOffice 和 pdftoppm")
    def test_real_conversion_reports_success_and_decodable_pages(self):
        source = self.base / "二次函数x².docx"
        doc = Document()
        doc.add_paragraph("Console conversion fixture 1+1=2.")
        doc.save(source)
        for encoding in ("gbk", "utf-8"):
            with self.subTest(encoding=encoding):
                output = self.base / encoding
                result = self.cli("doc_to_outputs.py", source, output, "72", encoding=encoding)
                self.assertEqual(result.returncode, 0, result.stderr)
                pdfs = list(output.rglob("*.pdf"))
                pages = list(output.rglob("*.png"))
                self.assertEqual(len(pdfs), 1)
                self.assertEqual(len(pages), 1)
                self.assertTrue(pdfs[0].read_bytes().startswith(b"%PDF-"))
                with Image.open(pages[0]) as image:
                    image.verify()
                with Image.open(pages[0]) as image:
                    image.load()
                    self.assertGreater(image.width, 0)
                    self.assertGreater(image.height, 0)
                output_text = result.stdout.decode(encoding, errors="replace")
                expected = "二次函数x²" if encoding == "utf-8" else "二次函数x\\xb2"
                self.assertIn(expected, output_text)


if __name__ == "__main__":
    unittest.main()
